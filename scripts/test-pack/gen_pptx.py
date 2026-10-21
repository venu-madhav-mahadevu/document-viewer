# PowerPoint test deck. usage: python gen_pptx.py <PREFIX> <output.pptx>   PREFIX = PPTX | PPT
import os, sys
from tp_common import ASSETS, Registry
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from lxml import etree

PREFIX, OUTFILE = sys.argv[1], sys.argv[2]
FILE = {'PPTX': 'test-slides.pptx', 'PPT': 'test-slides.ppt'}[PREFIX]
R = Registry(f'slides-{PREFIX.lower()}')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK, TITLE, BODY = prs.slide_layouts[6], prs.slide_layouts[0], prs.slide_layouts[1]
n = [0]

def panel(slide, slide_no, tests):
    """Gray 'tests on this slide' panel on the right."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.2), Inches(0.3), Inches(3.9), Inches(6.9))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2); box.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    for s in ('margin_left', 'margin_right', 'margin_top'):
        setattr(tf, s, Inches(0.12))
    p = tf.paragraphs[0]; r = p.add_run(); r.text = 'Tests on this slide'
    r.font.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    for what, expected, *status in tests:
        n[0] += 1
        tid = f'{PREFIX}-{n[0]:02d}'
        R.add(tid, FILE, f'Slide {slide_no}', what, expected, status[0] if status else '')
        p = tf.add_paragraph(); p.space_before = Pt(6)
        r = p.add_run(); r.text = f'{tid} · {what}'; r.font.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        p = tf.add_paragraph()
        r = p.add_run(); r.text = f'Expected: {expected}'; r.font.size = Pt(9); r.font.italic = True; r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

def title_box(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(8.6), Inches(0.8))
    r = tb.text_frame.paragraphs[0].add_run(); r.text = text
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

def set_bullet(p, char, font):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buNone', 'a:buChar', 'a:buFont', 'a:buAutoNum'):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    if font:
        f = etree.SubElement(pPr, qn('a:buFont')); f.set('typeface', font)
        if font in ('Wingdings', 'Symbol'): f.set('pitchFamily', '2'); f.set('charset', '2')
    b = etree.SubElement(pPr, qn('a:buChar')); b.set('char', char)

def set_autonum(p, kind):
    pPr = p._p.get_or_add_pPr()
    a = etree.SubElement(pPr, qn('a:buAutoNum')); a.set('type', kind)

# ---------------- 1. Title slide
s = prs.slides.add_slide(TITLE)
s.shapes.title.text = f'{FILE} — PowerPoint test deck'
s.placeholders[1].text = 'Each slide lists its tests on the right: ID, what is tested, expected result'
for sh in (s.shapes.title, s.placeholders[1]):
    sh.width = Inches(8.6); sh.left = Inches(0.4)
first = [('Title and subtitle', 'Large title text, smaller gray subtitle below it')]
if PREFIX == 'PPT':
    first.insert(0, ('PowerPoint 97-2003 file (.ppt)', 'Opens normally with all slides'))
first += [('Slide count and order', 'Page bar shows 10 slides; thumbnails in order 1-10'), ('16:9 slide size', 'Slides are wide (16:9), not square-ish (4:3)')]
panel(s, 1, first)

# ---------------- 2. Text formatting
s = prs.slides.add_slide(BLANK); title_box(s, '2. Text formatting')
tb = s.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(8.6), Inches(5.5)); tf = tb.text_frame; tf.word_wrap = True
lines = [
    ('Bold text', {'bold': True}), ('Italic text', {'italic': True}), ('Underlined text', {'underline': True}), ('Strikethrough text', {'strike': True}),
    ('Red text', {'color': 'C00000'}), ('Blue text', {'color': '0070C0'}), ('Green text', {'color': '00B050'}),
    ('Small 12 pt', {'size': 12}), ('Large 32 pt', {'size': 32}),
    ('Font Times New Roman', {'font': 'Times New Roman'}), ('Font Courier New', {'font': 'Courier New'}), ('Font Georgia', {'font': 'Georgia'}),
]
for i, (text, f) in enumerate(lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r = p.add_run(); r.text = text; r.font.size = Pt(f.get('size', 20))
    if f.get('bold'): r.font.bold = True
    if f.get('italic'): r.font.italic = True
    if f.get('underline'): r.font.underline = True
    if f.get('strike'): r._r.get_or_add_rPr().set('strike', 'sngStrike')
    if f.get('color'): r.font.color.rgb = RGBColor.from_string(f['color'])
    if f.get('font'): r.font.name = f['font']
p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = 'Centered line with superscript x'; r.font.size = Pt(20)
r = p.add_run(); r.text = '2'; r.font.size = Pt(20); r._r.get_or_add_rPr().set('baseline', '30000')
panel(s, 2, [('Bold, italic, underline, strikethrough', 'Lines 1-4 look like their names'), ('Text colors', 'Red, blue and green lines in those colors'),
             ('Font sizes', '"Small 12 pt" is small, "Large 32 pt" is big'), ('Font families', 'Times, Courier and Georgia lines in those fonts'),
             ('Center alignment and superscript', 'Last line is centered; the 2 is small and raised')])

# ---------------- 3. Bullets
s = prs.slides.add_slide(BODY)
s.shapes.title.text = '3. Bullets'; s.shapes.title.left = Inches(0.4); s.shapes.title.width = Inches(8.6)
body = s.placeholders[1]; body.left, body.top, body.width, body.height = Inches(0.4), Inches(1.4), Inches(8.6), Inches(5.8)
tf = body.text_frame
items = [(0, 'Round bullet • (Arial)', '•', 'Arial'), (0, 'Square bullet (Wingdings §)', '§', 'Wingdings'), (0, 'Box bullet (Wingdings q)', 'q', 'Wingdings'),
         (0, 'Filled square (Wingdings n)', 'n', 'Wingdings'), (0, 'Arrow bullet (Wingdings Ø)', 'Ø', 'Wingdings'), (0, 'Check bullet (Wingdings ü)', 'ü', 'Wingdings'),
         (0, 'Diamond bullet (Wingdings v)', 'v', 'Wingdings'), (0, 'Level 1', '•', 'Arial'), (1, 'Level 2 (indented more)', '–', 'Arial'),
         (2, 'Level 3 (square)', '§', 'Wingdings'), (3, 'Level 4 (smaller still)', '•', 'Arial')]
for i, (lvl, text, char, font) in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = text; p.level = lvl; p.runs[0].font.size = Pt(18 - 2 * lvl)
    set_bullet(p, char, font)
for kind, text in [('arabicPeriod', 'Numbered 1.'), ('arabicPeriod', 'Numbered 2.'), ('alphaLcParenR', 'Letter a)')]:
    p = tf.add_paragraph(); p.text = text; p.runs[0].font.size = Pt(18); set_autonum(p, kind)
panel(s, 3, [('Round bullet', 'Round •'), ('Wingdings §', 'Small square ▪ — NOT §'), ('Wingdings q', 'Hollow box ❑ — NOT q'), ('Wingdings n', 'Filled square ■ — NOT n'),
             ('Wingdings Ø', 'Arrow head ➢ — NOT Ø'), ('Wingdings ü', 'Check ✓ — NOT ü'), ('Wingdings v', 'Diamond ❖ — NOT v'),
             ('Multi-level indents', 'Levels 1-4 are indented further each time and get smaller'), ('Numbered and lettered', '1. 2. then a)')])

# ---------------- 4. Shrink text on overflow
s = prs.slides.add_slide(BODY)
s.shapes.title.text = '4. Shrink text on overflow'; s.shapes.title.left = Inches(0.4); s.shapes.title.width = Inches(8.6)
body = s.placeholders[1]; body.left, body.top, body.width, body.height = Inches(0.4), Inches(1.4), Inches(8.6), Inches(4.2)
tf = body.text_frame
for i in range(14):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = f'Line {i + 1}: this placeholder has more text than fits, so PowerPoint shrinks it.'; p.level = 1 if i % 3 else 0
bodyPr = tf._txBody.find(qn('a:bodyPr'))
for c in list(bodyPr): bodyPr.remove(c)
fit = etree.SubElement(bodyPr, qn('a:normAutofit')); fit.set('fontScale', '55000'); fit.set('lnSpcReduction', '20000')
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(5.65), Inches(8.6), Inches(0.05))
line.fill.solid(); line.fill.fore_color.rgb = RGBColor(0xC0, 0x00, 0x00); line.line.fill.background()
panel(s, 4, [('Shrink text on overflow', 'All 14 lines fit above the red line (text is smaller); nothing runs past it')])

# ---------------- 5. Shapes
s = prs.slides.add_slide(BLANK); title_box(s, '5. Shapes and arrows')
shape_list = [(MSO_SHAPE.RECTANGLE, 'Rectangle', '4472C4'), (MSO_SHAPE.ROUNDED_RECTANGLE, 'Rounded', '70AD47'), (MSO_SHAPE.OVAL, 'Circle', 'ED7D31'),
              (MSO_SHAPE.ISOSCELES_TRIANGLE, 'Triangle', 'FFC000'), (MSO_SHAPE.RIGHT_ARROW, 'Right arrow', 'C00000'), (MSO_SHAPE.LEFT_RIGHT_ARROW, 'Both ways', '7030A0'),
              (MSO_SHAPE.UP_ARROW, 'Up', '00B0F0'), (MSO_SHAPE.STAR_5_POINT, 'Star', 'FFD966'), (MSO_SHAPE.RECTANGULAR_CALLOUT, 'Callout', 'A9D18E'),
              (MSO_SHAPE.CHEVRON, 'Chevron', '8FAADC'), (MSO_SHAPE.PENTAGON, 'Pentagon', 'F4B183'), (MSO_SHAPE.HEXAGON, 'Hexagon', 'BF9000')]
for i, (kind, text, color) in enumerate(shape_list):
    x = Inches(0.4 + (i % 4) * 2.15); y = Inches(1.3 + (i // 4) * 1.55)
    sh = s.shapes.add_shape(kind, x, y, Inches(1.9), Inches(1.2))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(color); sh.line.color.rgb = RGBColor(0x40, 0x40, 0x40)
    sh.text_frame.text = text; sh.text_frame.paragraphs[0].runs[0].font.size = Pt(12); sh.text_frame.paragraphs[0].runs[0].font.bold = True
    sh.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
rot = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.1), Inches(1.6), Inches(0.7)); rot.rotation = 30
rot.fill.solid(); rot.fill.fore_color.rgb = RGBColor(0x5B, 0x9B, 0xD5); rot.text_frame.text = 'Rotated 30°'
outline = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.8), Inches(6.0), Inches(1.4), Inches(0.9)); outline.fill.background(); outline.line.color.rgb = RGBColor(0xC0, 0, 0); outline.line.width = Pt(3)
conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.6), Inches(6.45), Inches(8.6), Inches(6.45))
conn.line.color.rgb = RGBColor(0x00, 0x70, 0xC0); conn.line.width = Pt(3)
ln = conn.line._get_or_add_ln(); tail = etree.SubElement(ln, qn('a:tailEnd')); tail.set('type', 'triangle')
panel(s, 5, [('Basic shapes', 'Rectangle, rounded box, circle, triangle with their colors and labels'),
             ('Arrows', 'Right arrow, left-right arrow and up arrow have arrow shapes (not plain boxes)', '⚠️'),
             ('Other shapes', 'Star, callout, chevron, pentagon, hexagon have their real outlines', '⚠️'),
             ('Rotation', '"Rotated 30°" box is tilted'), ('Outline only', 'Red circle outline with no fill'),
             ('Connector with arrow head', 'Blue line at the bottom with an arrow head on the right', '⚠️')])

# ---------------- 6. Pictures
s = prs.slides.add_slide(BLANK); title_box(s, '6. Pictures')
s.shapes.add_picture(os.path.join(ASSETS, 'picture.png'), Inches(0.4), Inches(1.3), width=Inches(4.2))
crop = s.shapes.add_picture(os.path.join(ASSETS, 'picture.png'), Inches(4.8), Inches(1.3), width=Inches(4.2))
crop.crop_left, crop.crop_right, crop.crop_top, crop.crop_bottom = 0.0, 0.5, 0.0, 0.5
bord = s.shapes.add_picture(os.path.join(ASSETS, 'picture2.png'), Inches(0.4), Inches(4.4), width=Inches(3.6))
bord.line.color.rgb = RGBColor(0xC0, 0x00, 0x00); bord.line.width = Pt(4)
panel(s, 6, [('Picture', 'Top-left: whole "Embedded picture" test chart'), ('Cropped picture', 'Top-right: only the top-left quarter of the chart (title + RED/GREEN bars), enlarged', '⚠️'),
             ('Picture with border', 'Bottom: purple "Picture 2" with a thick red border', '⚠️')])

# ---------------- 7. Table
s = prs.slides.add_slide(BLANK); title_box(s, '7. Table')
tbl = s.shapes.add_table(5, 4, Inches(0.4), Inches(1.4), Inches(8.6), Inches(3)).table
data = [['Region', 'Q1', 'Q2', 'Total'], ['West', '1,200', '1,350', '2,550'], ['East', '1,300', '1,400', '2,700'], ['North', '1,400', '1,450', '2,850'], ['Merged note across all columns', '', '', '']]
for r_, row in enumerate(data):
    for c_, v in enumerate(row):
        cell = tbl.cell(r_, c_); cell.text = v
        if v: cell.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
tbl.cell(4, 0).merge(tbl.cell(4, 3))
tbl.cell(2, 3).fill.solid(); tbl.cell(2, 3).fill.fore_color.rgb = RGBColor(0xC6, 0xEF, 0xCE)
panel(s, 7, [('Table header', 'First row has a colored header with white bold text (theme style)'), ('Banded rows', 'Rows alternate light / darker shading'),
             ('Cell color', 'East total (2,700) has a light green background'), ('Merged cells', 'Last row is one cell across all 4 columns')])

# ---------------- 8. Groups, gradient, transparency
s = prs.slides.add_slide(BLANK); title_box(s, '8. Group, gradient, transparency')
grp = s.shapes.add_group_shape()
for i, color in enumerate(['4472C4', 'ED7D31', '70AD47']):
    g = grp.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + i * 1.4), Inches(1.6 + i * 0.3), Inches(1.6), Inches(1.0))
    g.fill.solid(); g.fill.fore_color.rgb = RGBColor.from_string(color); g.text_frame.text = f'Group {i + 1}'
grad = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(3.8), Inches(4), Inches(1.4))
grad.fill.gradient(); grad.fill.gradient_angle = 0
stops = grad.fill.gradient_stops; stops[0].color.rgb = RGBColor(0xFF, 0xFF, 0xFF); stops[1].color.rgb = RGBColor(0x1F, 0x4E, 0x79)
grad.text_frame.text = 'Gradient white -> dark blue'
tr = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.2), Inches(3.6), Inches(2), Inches(2))
tr.fill.solid(); tr.fill.fore_color.rgb = RGBColor(0xC0, 0x00, 0x00)
sf = tr.fill._xPr.find(qn('a:solidFill'))[0]; alpha = etree.SubElement(sf, qn('a:alpha')); alpha.set('val', '40000')
tr.text_frame.text = '40% opaque'
under = s.shapes.add_textbox(Inches(5.3), Inches(4.4), Inches(1.8), Inches(0.5)); under.text_frame.text = 'TEXT BEHIND'
s.shapes._spTree.remove(under._element); s.shapes._spTree.insert(2, under._element)
panel(s, 8, [('Grouped shapes', 'Three overlapping colored boxes "Group 1-3", stepping down to the right'),
             ('Gradient fill', 'Wide box fades from white (left) to dark blue (right)', '⚠️'),
             ('Transparency', 'Red circle is see-through: "TEXT BEHIND" is visible through it', '⚠️')])

# ---------------- 9. Charts
s = prs.slides.add_slide(BLANK); title_box(s, '9. Charts')
cd = CategoryChartData(); cd.categories = ['Q1', 'Q2', 'Q3', 'Q4']; cd.add_series('West', (12, 15, 18, 21)); cd.add_series('East', (10, 9, 14, 16))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.4), Inches(1.3), Inches(4.3), Inches(3.4), cd).chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM
pd = CategoryChartData(); pd.categories = ['Deposits', 'Loans', 'Other']; pd.add_series('Mix', (55, 35, 10))
pie = s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(4.8), Inches(1.3), Inches(4.2), Inches(3.4), pd).chart
pie.has_legend = True; pie.legend.position = XL_LEGEND_POSITION.BOTTOM
panel(s, 9, [('Column chart', 'Blue and orange columns for Q1-Q4 with a legend (West, East)', '❌'), ('Pie chart', 'Pie with 3 slices: Deposits 55%, Loans 35%, Other 10%', '❌')])

# ---------------- 10. Background and notes
s = prs.slides.add_slide(BLANK); title_box(s, '10. Slide background and speaker notes')
s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0xFF, 0xF2, 0xCC)
s.notes_slide.notes_text_frame.text = 'These are the speaker notes for slide 10.'
panel(s, 10, [('Slide background color', 'This slide has a light yellow background (other slides are white)'),
              ('Speaker notes', 'Notes "These are the speaker notes for slide 10." are available (if the viewer shows notes)', '❌')])

prs.save(OUTFILE)
R.save()
